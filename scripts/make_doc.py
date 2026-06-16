# -*- coding: utf-8 -*-
"""MAMOGRAF (AI Scan) loyihasi bo'yicha Word hujjat (.docx) generatori.

Diagrammalar (doc_assets/*.png) avval `make_diagrams.py` orqali yaratilishi kerak.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "MAMOGRAF_loyiha_malumot.docx"
ASSETS = ROOT / "doc_assets"

ACCENT = RGBColor(0x1F, 0x4E, 0x79)   # to'q ko'k
MUTED = RGBColor(0x55, 0x55, 0x55)


def setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for lvl, sz in ((1, 18), (2, 14), (3, 12)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = "Calibri"
        st.font.size = Pt(sz)
        st.font.color.rgb = ACCENT
        st.font.bold = True


def add_cover(doc: Document) -> None:
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("AI Scan — MAMOGRAF")
    r.bold = True
    r.font.size = Pt(30)
    r.font.color.rgb = ACCENT

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = s.add_run("Sun'iy intellekt asosida ko'krak bezi o'smalarini erta\n"
                   "aniqlash va diagnostika qilish tizimi")
    rs.font.size = Pt(14)
    rs.font.color.rgb = MUTED

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rm = meta.add_run(
        "Texnik-axborot hujjati (loyiha sharhi)\n"
        f"Tayyorlangan sana: {date.today().strftime('%d.%m.%Y')}\n"
        "Joriy ishlanma: multi-label + radiomics + WBF ansambl"
    )
    rm.font.size = Pt(11)
    rm.font.color.rgb = MUTED
    doc.add_page_break()


def h(doc, text, level=1):
    doc.add_heading(text, level=level)


def p(doc, text):
    return doc.add_paragraph(text)


def bullet(doc, text):
    return doc.add_paragraph(text, style="List Bullet")


def numbered(doc, text):
    return doc.add_paragraph(text, style="List Number")


def image(doc, name, width_in=6.3, caption=None):
    path = ASSETS / name
    if not path.exists():
        p(doc, f"[Diagramma topilmadi: {name}]")
        return
    pa = doc.add_paragraph()
    pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pa.add_run().add_picture(str(path), width=Inches(width_in))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cap.add_run(caption)
        cr.italic = True
        cr.font.size = Pt(9.5)
        cr.font.color.rgb = MUTED


def kv_table(doc, rows, headers=("Parametr", "Tavsif")):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = "Light Grid Accent 1"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(htext)
        run.bold = True
    for a, b in rows:
        c = tbl.add_row().cells
        c[0].text = str(a)
        c[1].text = str(b)
    doc.add_paragraph()
    return tbl


def build() -> None:
    doc = Document()
    setup_styles(doc)
    add_cover(doc)

    # 1. Umumiy ma'lumot
    h(doc, "1. Loyiha haqida umumiy ma'lumot")
    p(doc,
      "AI Scan (ichki nom — MAMOGRAF) — bu mammografiya tasvirlari va radiomika "
      "yondashuvlari yordamida ko'krak bezi o'smalarini (xavfsiz va xavfli) erta "
      "aniqlash hamda diagnostika qilishga mo'ljallangan sun'iy intellektga "
      "asoslangan tibbiy tizim. Loyiha ikki qismdan iborat: (1) anonimlashtirilgan "
      "mammografiya ma'lumotlar to'plami (dataset) va (2) to'liq veb-ilova "
      "(FastAPI backend + DICOM ko'ruvchi UI + ilmiy/o'qitish skriptlari + model "
      "arxitekturalari).")
    p(doc, "Loyiha quyidagilar uchun mo'ljallangan:")
    for x in ("Ko'krak bezi saratonini aniqlash (benign va malignant)",
              "BI-RADS klassifikatsiyasi",
              "Radiomika asosidagi tahlil",
              "Chuqur o'qitish (CNN, gibrid modellar) modellarini o'qitish",
              "Sun'iy intellektni klinik jarayonga (PACS bilan mos) integratsiya qilish"):
        bullet(doc, x)

    # 2. Maqsad va vazifalar
    h(doc, "2. Maqsad va asosiy vazifalar")
    for x in ("Ko'krak bezi saratonini erta bosqichda aniqlash",
              "Radiologlar ish jarayonini tezlashtirish va xatolarni kamaytirish",
              "Radiomika + chuqur o'qitish modellarini birgalikda qo'llash",
              "Diagnostika aniqligini oshirish va tahlil vaqtini qisqartirish",
              "AI ni klinik (PACS-mos) jarayonga xavfsiz integratsiya qilish"):
        bullet(doc, x)

    # 3. Tizim arxitekturasi (DIAGRAMMA)
    h(doc, "3. Tizim arxitekturasi")
    p(doc, "Tizim to'rt qatlamdan iborat: frontend (UI), FastAPI backend, AI/qayta "
           "ishlash qatlami va ma'lumot/integratsiya qatlami. Quyidagi diagramma "
           "qatlamlar va ular orasidagi bog'lanishni ko'rsatadi.")
    image(doc, "architecture.png", width_in=6.4,
          caption="1-rasm. MAMOGRAF tizimining qatlamli arxitekturasi")
    for a, b in [
        ("Frontend", "DICOM ko'ruvchi, annotatsiya muharriri, AI taklif ko'rinishi, Model Studio (o'qitish paneli), PACS/Worklist."),
        ("Backend (FastAPI)", "REST API + WebSocket: upload, annotations, training, inference va autentifikatsiya endpointlari."),
        ("AI qatlam", "BCA-YOLO, TILLNet-Det, XS-Classifier model arxitekturalari + Radiomics + Inference (WBF/TTA)."),
        ("Ma'lumot/integratsiya", "uploads/ (DICOM), annotations/ (JSON), SQLite DB, de-identify, PACS ulanishi, eksport modullari."),
    ]:
        bullet(doc, f"{a} — {b}")

    # 4. TO'LIQ ISH TSIKLI (asosiy bo'lim)
    h(doc, "4. To'liq ish tsikli: Upload → Auto-AI → Annotation → Dataset → Yangi model")
    p(doc, "MAMOGRAF yopiq (closed-loop) ish tsikliga asoslangan: model qancha "
           "ko'p ishlatilsa, shuncha ko'p tekshirilgan annotatsiya to'planadi, bu esa "
           "keyingi (yanada aniq) modelni o'qitish uchun ma'lumot bo'ladi. Bu yondashuv "
           "AKTIV O'QITISH (Active Learning) deb ataladi.")
    image(doc, "cycle.png", width_in=6.5,
          caption="2-rasm. To'liq ish tsikli — beshta bosqich va aktiv o'qitish qaytishi")

    h(doc, "4.1. Bosqich 1 — Upload (DICOM yuklash)", level=2)
    p(doc, "Endpoint: POST /api/upload. Foydalanuvchi bir yoki bir nechta DICOM faylni "
           "yuklaydi. Har bir fayl bilan quyidagilar bajariladi:")
    for x in ("Fayl noyob ID (uuid) bilan uploads/<id>.dcm sifatida saqlanadi",
              "pydicom bilan o'qib validatsiya qilinadi; yaroqsiz fayl rad etiladi (HTTP 400)",
              "AUTO_DEIDENTIFY yoqilgan bo'lsa — PHI tag'lari DARHOL tozalanadi (anonymize_in_place), ya'ni annotatsiya/eksportdan OLDIN",
              "Metama'lumot qaytariladi: o'lcham (rows/cols), modality, ViewPosition, ImageLaterality, kadrlar soni",
              "AUTO_INFER_ON_UPLOAD yoqilgan bo'lsa — AI inference fon (background) vazifasi sifatida navbatga qo'yiladi"):
        bullet(doc, x)
    p(doc, "Muhim: AI inference upload so'rovini bloklamaydi — foydalanuvchi DICOM'ni "
           "darhol ko'radi, AI belgilari (bbox) esa ~10–20 soniyada fon rejimida paydo bo'ladi.")

    h(doc, "4.2. Bosqich 2 — Auto-AI (avtomatik inference)", level=2)
    p(doc, "Funksiya: _auto_infer_uploaded. Yuklangan tasvir uchun YOLO modeli o'choqlarni "
           "(lesion) avtomatik aniqlaydi va pseudo-annotatsiyalar yaratadi:")
    for x in ("Mavjud modellardan biri tanlanadi (AUTO_INFER_MODEL yoki birinchi mavjud model)",
              "DICOM → PNG ga render qilinadi va inf.infer_png(...) chaqiriladi",
              "Har bir deteksiya ishonch (confidence) qiymati bo'yicha 3 zonaga ajratiladi",
              "Natijalar created_by='ai:<model>' va status (ai_accepted/ai_review/ai_suspect) bilan saqlanadi",
              "Funksiya hech qachon xato (exception) qaytarmaydi — upload buzilmasligi kafolatlanadi"):
        bullet(doc, x)
    image(doc, "zones.png", width_in=6.1,
          caption="3-rasm. AI ishonch (confidence) zonalari — avtomatik tasniflash")
    p(doc, "Ishonch < 0.20 bo'lgan deteksiyalar umuman saqlanmaydi. Bu chegaralar "
           "AUTO_INFER_ACCEPT_THR / REVIEW_THR / SUSPECT_THR muhit o'zgaruvchilari orqali sozlanadi.")

    h(doc, "4.3. Bosqich 3 — Annotation (radiolog tekshiruvi)", level=2)
    p(doc, "Radiolog interfeysda DICOM'ni ochganda AI takliflari rangli zonalar bilan "
           "allaqachon chizilgan bo'ladi. Radiolog quyidagilarni bajaradi:")
    for x in ("AI belgilarini tasdiqlaydi, tahrirlaydi yoki o'chiradi",
              "Zarur bo'lsa yangi bbox yoki polygon qo'lda qo'shadi",
              "Multi-label: bitta belgi bir nechta yorliqqa ega bo'lishi mumkin (labels: [...])",
              "Status o'zgaradi: edited / approved / submitted va h.k.",
              "Har bir o'zgarish audit jurnaliga yoziladi (kim, qachon, nima)"):
        bullet(doc, x)
    p(doc, "Endpointlar: PUT /api/annotations (saqlash), POST /api/annotations/status "
           "(holatni o'zgartirish), GET /api/annotations/history (tarix/audit). "
           "GET /api/training/suggestion esa tekshirilgan annotatsiyalar sonini sanab, "
           "ular ACTIVE_LEARNING_THRESHOLD (sukut bo'yicha 50) ga yetganda 'qayta "
           "o'qitishga vaqt keldi' degan tavsiya beradi.")

    h(doc, "4.4. Bosqich 4 — Training dataset (o'qitish to'plami)", level=2)
    p(doc, "Endpoint: POST /api/training/prepare. Annotatsiyalardan Ultralytics YOLO "
           "formatidagi dataset yaratiladi:")
    for x in ("annotations/upload__*.json fayllar o'qiladi",
              "Filtrlar qo'llanadi: include_ai (AI belgilarini qo'shish/qo'shmaslik), statuses",
              "bbox va polygon → YOLO normalized formatga aylantiriladi (polygon o'rab oluvchi bbox'ga)",
              "Classlar ro'yxati: berilgan class_list yoki yorliqlardan avtomatik",
              "Bemor darajasida (patient-level) train/val ajratish — bir bemorning rasmlari faqat bitta to'plamga (data leakage'ni oldini olish)",
              "seed orqali takrorlanuvchi (reproducible) bo'linish"):
        bullet(doc, x)
    p(doc, "Natija struktura:")
    code = doc.add_paragraph()
    cr = code.add_run("<dest>/\n  images/{train,val}/<id>.png\n  labels/{train,val}/<id>.txt\n  data.yaml")
    cr.font.name = "Consolas"
    cr.font.size = Pt(10)

    h(doc, "4.5. Bosqich 5 — Yangi model (o'qitish va deploy)", level=2)
    p(doc, "Endpoint: POST /api/training/run (fonda _run_yolo_train). data.yaml "
           "tekshiriladi va Ultralytics YOLO o'qitish fon rejimida boshlanadi:")
    for x in ("base_model (masalan yolo11n.pt), epochs, imgsz, batch — asosiy parametrlar",
              "optimizer, lr0/lrf, momentum, weight_decay, warmup, patience (early stopping), cos_lr",
              "Augmentation: hsv, fliplr/flipud, scale, mosaic, mixup",
              "device (CPU/GPU), workers, cache — apparat sozlamalari",
              "deploy_after=True bo'lsa — o'qitilgan model avtomatik joylanadi (deploy)",
              "Holat kuzatuvi: GET /api/training/status/{run_id}, GET /api/training/runs"):
        bullet(doc, x)
    p(doc, "Yangi model joylangach, u keyingi Auto-AI bosqichida ishlatiladi — shu tariqa "
           "sikl yopiladi va har bir aylanishda model aniqligi oshib boradi.")

    # 5. Asosiy imkoniyatlar
    h(doc, "5. Tizimning asosiy imkoniyatlari")
    h(doc, "5.1. Veb-ilova (FastAPI)", level=2)
    for x in ("DICOM fayllarni yuklash, ko'rish va kadrlar bo'yicha ko'rib chiqish",
              "Anatomik joylashuvni belgilash (bounding box / segmentatsiya annotatsiyasi)",
              "Multi-label annotatsiya — bitta belgi bir nechta yorliqqa ega bo'lishi mumkin",
              "AI taklif (suggestion) — model aniqlagan o'choqlarni avtomatik chizish",
              "Ishonch (confidence) zonalari bo'yicha 3 darajali tasniflash",
              "Annotatsiyalar tarixi, audit jurnali va bildirishnomalar",
              "JWT + TOTP (ikki faktorli) autentifikatsiya va foydalanuvchilar boshqaruvi",
              "Real vaqt yangilanishlari uchun WebSocket kanali",
              "Veb-interfeysda 'Model Studio' — modelni o'qitishni boshqarish paneli (train.html)"):
        bullet(doc, x)

    h(doc, "5.2. PACS integratsiyasi", level=2)
    for x in ("PACS serverlarini ro'yxatga olish va C-ECHO bilan ulanishni tekshirish",
              "C-FIND (query), C-STORE (store) va C-GET/C-MOVE (fetch) operatsiyalari",
              "Modality Worklist (MWL) — ish ro'yxatini import qilish",
              "PACS dan kelgan tadqiqotlarni ish ro'yxatiga (worklist) qo'shish"):
        bullet(doc, x)

    h(doc, "5.3. Eksport formatlari", level=2)
    for x in ("YOLO detection dataseti (.zip — labels + classes + data.yaml)",
              "Pascal VOC XML dataseti",
              "CSV jadval (annotatsiyalar va metama'lumotlar)",
              "DICOM-SEG (segmentatsiya niqobi) va DICOM-SR (struktura hisoboti)",
              "Niqob (mask) PNG yoki NIfTI formatida",
              "Radiomika belgilarini CSV ga eksport qilish"):
        bullet(doc, x)

    # 6. AI model arxitekturalari
    h(doc, "6. AI model arxitekturalari (app/models_arch/)")
    p(doc, "Loyiha tarkibida uchta original model arxitekturasi mavjud (vaznlarsiz, "
           "ya'ni faqat kod/struktura sifatida):")

    h(doc, "6.1. BCA-YOLO", level=2)
    p(doc, "Bilateral Cross-Attention YOLO — mammografik o'choqlarni aniqlash uchun. "
           "Chap va o'ng ko'krak tasvirlari orasidagi assimetriyani 'cross-attention' "
           "mexanizmi orqali hisobga oladi:")
    for x in ("Umumiy YOLO backbone — ko'p masshtabli xususiyatlar (F_v)",
              "Bilateral Cross-Attention (BCA) — chap/o'ng proyeksiyalarni juftlash",
              "Inter-View Consistency (IVC) — CC va MLO ko'rinishlarini birlashtirish",
              "Deteksiya boshlari (har bir ko'rinish + birlashtirilgan yordamchi)",
              "Ordinal BI-RADS regressiya boshi (cumulative link)"):
        bullet(doc, x)

    h(doc, "6.2. TILLNet-Det", level=2)
    p(doc, "Text-Informed Lesion Localisation Network — radiologiya hisobotini (matnni) "
           "vizual xususiyatlar piramidasiga shartlashtiruvchi multimodal detektor:")
    for x in ("Cross-script matn enkoderi — Kirill va Lotin (o'zbek-Cy/o'zbek-Lat/rus/ingliz) uchun belgi darajasidagi Transformer",
              "FiLM-fused FPN — matn vektori bilan xususiyatlar piramidasini modulyatsiya qilish",
              "Anchor-free FCOS uslubidagi bosh — klassifikatsiya, regressiya va centerness",
              "Boshlar piramida darajalari bo'ylab umumiy (~1.6M parametr — kichik korpusga mos)"):
        bullet(doc, x)

    h(doc, "6.3. XS-Classifier", level=2)
    p(doc, "Cross-Script Classifier — ko'p tilli (o'zbek-Kirill / o'zbek-Lotin / rus) "
           "radiologiya hisobotlarini tasniflash uchun TF-IDF + Logistic Regression "
           "asosidagi yengil model.")

    # 7. Radiomika
    h(doc, "7. Radiomika moduli (app/radiomics.py)")
    p(doc, "Annotatsiya qilingan ROI (qiziqish hududi) ichidan miqdoriy belgilarni "
           "ajratib oladi. Quyidagi belgi oilalari hisoblanadi:")
    for x in ("First-order statistik belgilar (o'rtacha, dispersiya, entropiya va h.k.)",
              "Shape (shakl) belgilari — yuza, perimetr, kompaktlik",
              "GLCM — Gray-Level Co-occurrence Matrix",
              "GLRLM — Gray-Level Run-Length Matrix",
              "GLSZM — Gray-Level Size-Zone Matrix",
              "NGTDM — Neighbouring Gray-Tone Difference Matrix"):
        bullet(doc, x)

    # 8. Ma'lumotlar maxfiyligi
    h(doc, "8. Ma'lumotlar maxfiyligi va de-identifikatsiya")
    p(doc, "Tizim har bir yuklangan DICOM fayldan PHI (shaxsiy identifikatsiya "
           "ma'lumotlari) ni annotatsiya yoki eksportdan OLDIN avtomatik olib tashlaydi. "
           "Keyingi barcha bosqichlar faqat anonimlashtirilgan nusxa ustida ishlaydi.")
    for x in ("PatientName → ANONYMOUS^PATIENT",
              "PatientID → ANON",
              "ReferringPhysicianName, OperatorsName, InstitutionName, StationName, AccessionNumber → tozalanadi",
              "StudyDate/AcquisitionDate → yil saqlanadi, oy/kun 0101 ga o'rnatiladi",
              "PatientIdentityRemoved tegi → YES",
              "Piksel ma'lumotlari, o'lcham va Modality SAQLANADI (AI/tadqiqot uchun)"):
        bullet(doc, x)
    p(doc, "Sozlama: AUTO_DEIDENTIFY=0 o'rnatilsa o'chiriladi (sukut bo'yicha yoqilgan). "
           "Mavjud fayllarni qayta anonimlashtirish: python -m app.deidentify app/uploads")

    # 9. Ma'lumotlar to'plami
    h(doc, "9. Ma'lumotlar to'plami (dataset)")
    p(doc, "Ma'lumotlar 'encrypted_dataset_grouped/' ichida GROUP_XXXXX papkalari "
           "ko'rinishida saqlanadi — har bir guruh bitta tadqiqot (bemor holati). "
           "Metama'lumotlar mapping_clean.xlsx faylida quyidagi maydonlar bilan:")
    kv_table(doc, [
        ("study_id / group_code", "Anonim tadqiqot identifikatori"),
        ("image_code", "Tasvir ID"),
        ("image_index", "Tasvir tartibi"),
        ("view_position", "CC, MLO va h.k."),
        ("body_part", "Ko'krak (Breast)"),
        ("manufacturer / model", "Qurilma ishlab chiqaruvchisi / modeli"),
        ("presentation_type", "Tasvir formati / turi"),
        ("status", "Qayta ishlash holati"),
    ], headers=("Ustun", "Tavsif"))

    # 10. Texnologiyalar
    h(doc, "10. Texnologiyalar steki")
    kv_table(doc, [
        ("Backend", "Python, FastAPI, Uvicorn"),
        ("AI / ML", "PyTorch, Ultralytics (YOLO), scikit-learn, scikit-image, SciPy"),
        ("DICOM", "pydicom, pylibjpeg, highdicom, nibabel"),
        ("Tasvirlash", "NumPy, Pillow, OpenCV (headless)"),
        ("Xavfsizlik", "PyJWT, bcrypt, pyotp (TOTP), slowapi (rate limit)"),
        ("Frontend", "HTML / JavaScript / CSS (statik UI)"),
        ("Ma'lumotlar", "SQLite (DB), openpyxl (XLSX)"),
        ("Deploy", "Docker / docker compose"),
    ], headers=("Qatlam", "Texnologiyalar"))

    # 11. O'rnatish
    h(doc, "11. O'rnatish va ishga tushirish")
    numbered(doc, "Virtual muhit yaratish: python -m venv .venv")
    numbered(doc, "Faollashtirish (Windows): .venv\\Scripts\\activate")
    numbered(doc, "Kutubxonalarni o'rnatish: pip install -r requirements.txt")
    numbered(doc, "Ishga tushirish: python -m uvicorn app.main:app --reload")
    p(doc, "Yoki Docker orqali: docker compose up --build")
    p(doc, "Eslatma: o'qitilgan model vaznlari (*.pt) va bemor ma'lumotlari (DICOM, "
           "annotatsiyalar, SQLite DB) repozitoriyga kiritilmagan (.gitignore).")

    # 12. Joriy holat
    h(doc, "12. Loyihaning joriy holati")
    p(doc, "Joriy ishlanma (git branch: feature/multilabel-radiomics-ensemble) "
           "doirasida quyidagilar amalga oshirilgan:")
    for x in ("Multi-label annotatsiya qo'llab-quvvatlash",
              "Radiomika belgilarini ajratish va CSV eksporti",
              "WBF (Weighted Boxes Fusion) ansambl + TTA",
              "Yangi eksport formatlari (DICOM-SEG, DICOM-SR, NIfTI)",
              "Avtomatik DICOM anonimlashtirish (upload paytida)",
              "Veb-interfeysda Model Studio (o'qitish paneli — train.html/train.js)"):
        bullet(doc, x)

    # 13. Cheklovlar
    h(doc, "13. Cheklovlar va muhim eslatmalar")
    for x in ("Ma'lumotlar to'plami faqat ilmiy-tadqiqot maqsadlari uchun",
              "To'g'ridan-to'g'ri klinik foydalanish uchun mo'ljallanmagan",
              "Joriy etishdan oldin validatsiya talab qilinadi",
              "Etika va me'yoriy talablarga rioya qilish shart",
              "Litsenziya: notijorat ilmiy foydalanish uchun"):
        bullet(doc, x)

    doc.add_paragraph()
    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    er = end.add_run("— AI Scan / MAMOGRAF loyihasi doirasida tayyorlandi —")
    er.italic = True
    er.font.color.rgb = MUTED

    doc.save(str(OUT))
    print(f"Saqlandi: {OUT}")


if __name__ == "__main__":
    build()
