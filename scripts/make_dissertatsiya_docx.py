# -*- coding: utf-8 -*-
"""MAMOGRAF — Magistrlik dissertatsiyasi (.docx) generatori.

05.01.11 — Raqamli texnologiyalar va sun'iy intellekt yo'nalishi.
Namuna (Abdiyeva X.S.) uslubida: titul, mundarija, kirish, 4 bob (§lar +
bob xulosalari), xulosa, foydalanilgan adabiyotlar, ilovalar.
Formulalar matplotlib mathtext orqali PNG ga chiziladi va joylanadi.

Manba: BCA-YOLO va Radiomic Faster R-CNN maqolalari (muallifning o'z ishi) +
MAMOGRAF dasturiy majmuasi (ushbu repozitoriy).
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "doc_assets"
EQ = ASSETS / "eq"
EQ.mkdir(parents=True, exist_ok=True)
OUT = ROOT / "MAMOGRAF_magistrlik_dissertatsiya.docx"

ACCENT = RGBColor(0x14, 0x2C, 0x52)
MUTED = RGBColor(0x55, 0x55, 0x55)


# --------------------------------------------------------------------------- #
# Formula renderer (LaTeX -> PNG)                                             #
# --------------------------------------------------------------------------- #
def render_eq(name: str, latex: str, fontsize: int = 20) -> Path:
    path = EQ / f"{name}.png"
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.text(0.0, 0.0, f"${latex}$", fontsize=fontsize, color="black")
    fig.savefig(str(path), dpi=200, bbox_inches="tight", pad_inches=0.12,
                transparent=False, facecolor="white")
    plt.close(fig)
    return path


FORMULAS = {
    "bca_attn": r"A \;=\; \mathrm{softmax}\!\left(\frac{Q\,K^{\top}}{\sqrt{d}}\right),\qquad \Delta F_L = A\,V_R",
    "bca_res": r"F'_L \;=\; F_L \;+\; \alpha\,\cdot\,\Delta F_L,\qquad \alpha \leftarrow 0",
    "asym": r"\mathcal{L}_{asym} \;=\; \frac{1}{|\Omega|}\sum_{(s,v)\in\Omega}\left(\bar{A}^{(s)}_v - m^{(s)}_v\right)^2",
    "ivc_gap": r"\hat{o}_v \;=\; \frac{1}{H\,W}\sum_{i=1}^{H W} o_v(i)",
    "ivc": r"\mathcal{L}_{IVC} \;=\; \left(\hat{o}_{CC} - \hat{o}_{MLO}\right)^2",
    "ordinal": r"\mathcal{L}_{ord} = -\sum_{k=0}^{K-2}\left( t_k\,\log\sigma(\theta_k-f) + (1-t_k)\,\log(1-\sigma(\theta_k-f)) \right),\ \ t_k=\mathbf{1}(y\leq k)",
    "thresh": r"\theta_k \;=\; \theta_0 + \sum_{j=1}^{k}\mathrm{softplus}(\tau_j),\qquad \mathrm{softplus}(\tau)=\log(1+e^{\tau})",
    "grad_alpha": r"\frac{\partial \mathcal{L}}{\partial \alpha} \;=\; \left\langle \frac{\partial \mathcal{L}}{\partial F'},\ \Delta F \right\rangle_F",
    "total": r"\mathcal{L} = \mathcal{L}_{box} + \mathcal{L}_{cls} + \mathcal{L}_{dfl} + \lambda_1\mathcal{L}_{asym} + \lambda_2\mathcal{L}_{IVC} + \lambda_3\mathcal{L}_{ord}",
    "iou": r"\mathrm{IoU}(B_p,B_g) \;=\; \frac{|B_p \cap B_g|}{|B_p \cup B_g|}",
    "prf": r"P=\frac{TP}{TP+FP},\quad R=\frac{TP}{TP+FN},\quad F_1=\frac{2\,P\,R}{P+R}",
    "map": r"\mathrm{mAP} \;=\; \frac{1}{C}\sum_{c=1}^{C}\int_{0}^{1} P_c(R)\,dR",
    "dice": r"\mathrm{Dice}(X,Y) \;=\; \frac{2\,|X\cap Y|}{|X|+|Y|}",
    "glcm": r"\mathrm{Contrast}=\sum_{i,j}(i-j)^2\,p(i,j),\qquad H=-\sum_{i,j}p(i,j)\log p(i,j)",
    "fusion": r"z \;=\; \phi\left(\,[\,f_{deep}\,\Vert\, f_{radiomic}\,]\,\right),\qquad \hat{y}=\mathrm{softmax}(W z + b)",
}


# --------------------------------------------------------------------------- #
# Word helpers                                                                #
# --------------------------------------------------------------------------- #
def setup(doc):
    n = doc.styles["Normal"]
    n.font.name = "Times New Roman"
    n.font.size = Pt(14)
    pf = n.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)
    for lvl, sz in ((1, 15), (2, 14), (3, 14)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = "Times New Roman"
        st.font.size = Pt(sz)
        st.font.bold = True
        st.font.color.rgb = ACCENT


def h1(doc, t):
    p = doc.add_heading(level=1)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t.upper())
    r.bold = True


def h2(doc, t):
    doc.add_heading(t, level=2)


def para(doc, t, just=True, bold=False, italic=False, muted=False, first_line=True):
    p = doc.add_paragraph()
    if just:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_line:
        p.paragraph_format.first_line_indent = Inches(0.4)
    # **bold** segmentlarini ajratish
    for i, seg in enumerate(t.split("**")):
        r = p.add_run(seg)
        r.bold = bold or (i % 2 == 1)
        r.italic = italic
        if muted:
            r.font.color.rgb = MUTED
    return p


def lead(doc, label, text):
    """Qalin sarlavhali abzas (KIRISH bo'limlari uchun): 'Maqsad. matn...'"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.4)
    r = p.add_run(label + " ")
    r.bold = True
    p.add_run(text)
    return p


def bullets(doc, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for it in items:
        p = doc.add_paragraph(style=style)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for i, seg in enumerate(it.split("**")):
            r = p.add_run(seg)
            r.bold = (i % 2 == 1)


def eq(doc, name, number=None, width=None):
    path = EQ / f"{name}.png"
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    # kenglikni mazmuniga qarab moslash
    from PIL import Image
    iw, ih = Image.open(path).size
    w = width or min(6.0, max(2.0, iw / 220))
    run.add_picture(str(path), width=Inches(w))
    if number:
        tp = p.add_run("    (" + str(number) + ")")
        tp.font.size = Pt(12)


def img(doc, name, width=6.0, caption=None):
    path = ASSETS / name
    if not path.exists():
        para(doc, f"[rasm: {name}]", muted=True, italic=True, first_line=False)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(caption)
        r.italic = True
        r.font.size = Pt(12)


def table(doc, headers, rows, caption=None):
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(caption)
        r.bold = True
        r.font.size = Pt(12)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, hh in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        rr = cell.paragraphs[0].add_run(hh)
        rr.bold = True
        rr.font.size = Pt(12)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            rr = cells[i].paragraphs[0].add_run(str(val))
            rr.font.size = Pt(12)
    doc.add_paragraph()


def chapter_concl(doc, bob, items):
    h2(doc, f"{bob} bob bo'yicha xulosalar")
    bullets(doc, items, numbered=True)


# --------------------------------------------------------------------------- #
# Document                                                                     #
# --------------------------------------------------------------------------- #
def build():
    for nm, tex in FORMULAS.items():
        render_eq(nm, tex)
    print("formulalar chizildi:", len(FORMULAS))

    doc = Document()
    setup(doc)

    # ===== TITUL =====
    def center(txt, sz=14, bold=False, sp=0):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(sp)
        r = p.add_run(txt)
        r.bold = bold
        r.font.size = Pt(sz)
        return p

    center("OʻZBEKISTON RESPUBLIKASI OLIY TAʼLIM, FAN VA INNOVATSIYALAR VAZIRLIGI", 13, True, 2)
    center("[OLIY TAʼLIM MUASSASASI NOMI]", 13, True, 18)
    center("Qoʻlyozma huquqida", 12)
    center("UDK 004.93", 12, sp=24)
    center("[MUALLIF FAMILIYASI ISMI OTASINING ISMI]", 14, True, 18)
    center("MAMMOGRAFIYA TASVIRLARI ASOSIDA KOʻKRAK OʻSMA SOHALARINI "
           "ANIQLASH, BI-RADS BOʻYICHA TASNIFLASH VA AVTOMATIK XULOSA "
           "SHAKLLANTIRISHNING SUNʼIY INTELLEKTGA ASOSLANGAN DASTURIY MAJMUASI",
           14, True, 18)
    center("05.01.11 — Raqamli texnologiyalar va sunʼiy intellekt", 13, True, 6)
    center("Texnika fanlari magistri akademik darajasini olish uchun yozilgan", 12, sp=2)
    center("MAGISTRLIK DISSERTATSIYASI", 15, True, 30)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("Ilmiy rahbar: [Ilmiy rahbar F.I.Sh., ilmiy darajasi, unvoni]").font.size = Pt(12)
    center(f"Toshkent — {date.today().year}", 13, True, 0)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nr = note.add_run("[Eslatma: titul varaqdagi [...] joylarni — muassasa, muallif, "
                      "ilmiy rahbar — o'zingiz to'ldiring.]")
    nr.italic = True
    nr.font.size = Pt(10)
    nr.font.color.rgb = MUTED
    doc.add_page_break()

    # ===== MUNDARIJA =====
    h1(doc, "Mundarija")
    toc = [
        "KIRISH",
        "I BOB. MAMMOGRAFIYA TASVIRLARINI TAHLIL QILISHNING ZAMONAVIY HOLATI",
        "  1.1-§. Koʻkrak saratoni, skrining mammografiyasi va BI-RADS tizimi",
        "  1.2-§. Tibbiy tasvirlarga ishlov berish va DICOM standarti",
        "  1.3-§. Chuqur oʻrganish va radiomika asosida oʻsma aniqlash usullari",
        "  1.4-§. Tadqiqot masalasining qoʻyilishi",
        "  I bob boʻyicha xulosalar",
        "II BOB. KOʻKRAK OʻSMALARINI ANIQLASH VA TASNIFLASH ALGORITMLARI",
        "  2.1-§. BCA-YOLO arxitekturasi: ikki tomonlama oʻzaro eʼtibor va koʻrinishlararo muvofiqlik",
        "  2.2-§. Ordinal BI-RADS regressiyasi va kompozit yoʻqotish funksiyasi",
        "  2.3-§. Radiomik belgilarni chuqur oʻrganish bilan birlashtirish",
        "  II bob boʻyicha xulosalar",
        "III BOB. MAMOGRAF DASTURIY MAJMUASINI AMALGA OSHIRISH",
        "  3.1-§. Tizim arxitekturasi, maxfiylik va xavfsizlik",
        "  3.2-§. Annotatsiya, AI yordami va model oʻqitish (Model Studio)",
        "  3.3-§. Avtomatik xulosa shakllantirish va DICOM SR/PACS integratsiyasi",
        "  III bob boʻyicha xulosalar",
        "IV BOB. TAJRIBAVIY TADQIQOTLAR VA AMALIYOTDA QOʻLLASH",
        "  4.1-§. Maʼlumotlar toʻplami va baholash metodologiyasi",
        "  4.2-§. Tajriba natijalari va tahlili",
        "  4.3-§. Dasturiy majmuani amaliyotda qoʻllash",
        "  IV bob boʻyicha xulosalar",
        "XULOSA",
        "FOYDALANILGAN ADABIYOTLAR ROʻYXATI",
        "ILOVALAR",
    ]
    for t in toc:
        p = doc.add_paragraph(t)
        p.paragraph_format.line_spacing = 1.4
        if not t.startswith("  "):
            p.runs[0].bold = True
    doc.add_page_break()

    # ===== KIRISH =====
    h1(doc, "Kirish")
    lead(doc, "Dissertatsiya mavzusining dolzarbligi va zarurati.",
         "Jahonda ko'krak saratoni ayollar o'rtasida eng ko'p uchraydigan onkologik "
         "kasalliklardan biri bo'lib qolmoqda. Jahon sog'liqni saqlash tashkiloti (JSST) "
         "va Globocan ma'lumotlariga ko'ra, har yili millionlab yangi holatlar qayd "
         "etilib, yuz minglab ayol bu kasallikdan vafot etadi. Kasallikni erta bosqichda "
         "aniqlash yashash ko'rsatkichini sezilarli oshiradi, bunda skrining mammografiyasi "
         "eng samarali usul hisoblanadi. Biroq mammografiya tasvirlarini qo'lda talqin "
         "qilish radiologning tajribasiga bog'liq, kuzatuvchilararo o'zgaruvchanlik, "
         "charchoq va vaqt cheklovlari kabi muammolarga duch keladi. Shu sababli "
         "tibbiyot tasvirlarini sun'iy intellekt yordamida avtomatik tahlil qiluvchi, "
         "radiolog qaroriga yordam beruvchi dasturiy majmualar yaratish dolzarb "
         "ilmiy-amaliy masaladir.")
    para(doc,
         "O'zbekiston Respublikasida onkologik kasalliklarni erta aniqlash va o'lim "
         "ko'rsatkichini kamaytirish davlat siyosatining ustuvor yo'nalishlaridan biriga "
         "aylangan. «O'zbekiston-2030» strategiyasi va «Raqamli O'zbekiston-2030» "
         "strategiyasida raqamli texnologiyalar va sun'iy intellektni sog'liqni saqlash "
         "tizimiga keng joriy etish vazifalari belgilangan. Ushbu dissertatsiya tadqiqoti "
         "mazkur vazifalarni amalga oshirishga — mammografiya tasvirlaridan ko'krak o'sma "
         "sohalarini aniqlash, BI-RADS bo'yicha tasniflash va avtomatik xulosa "
         "shakllantirish dasturiy majmuasini yaratishga — qaratilgan.")

    lead(doc, "Muammoning o'rganilganlik darajasi.",
         "Mammografik shikastlanishlarni aniqlashda zamonaviy chuqur o'rganish modellari "
         "(YOLO, Faster R-CNN va boshqalar) keng qo'llanilmoqda. Biroq aksariyat "
         "yondashuvlar to'rtta standart ko'rinishni (L-CC, R-CC, L-MLO, R-MLO) mustaqil "
         "qayta ishlaydi va radiologlar foydalanadigan ikki tomonlama (kontralateral) "
         "asimmetriya hamda ko'rinishlararo muvofiqlik signallarini hisobga olmaydi. "
         "Bundan tashqari, BI-RADS shkalasi tartibli (ordinal) bo'lishiga qaramay, ko'p "
         "modellar uni nominal sinflar sifatida talqin qiladi. Ushbu kamchiliklar "
         "tadqiqotning ilmiy yo'nalishini belgilab berdi.")

    lead(doc, "Tadqiqotning maqsadi —",
         "mammografiya tasvirlarida ko'krak o'sma sohalarini aniqlash, BI-RADS bo'yicha "
         "tasniflash va avtomatik klinik xulosa shakllantirishning sun'iy intellektga "
         "asoslangan, PACS bilan mos, amaliyotda qo'llanishga yaroqli dasturiy majmuasini "
         "(MAMOGRAF) ishlab chiqish.")

    para(doc, "**Tadqiqotning vazifalari:**", first_line=False)
    bullets(doc, [
        "mammografiya tasvirlarini tahlil qilishning zamonaviy holatini va mavjud usullarni o'rganish;",
        "ikki tomonlama o'zaro e'tibor (BCA) va ko'rinishlararo muvofiqlik (IVC) asosida o'sma aniqlash algoritmini (BCA-YOLO) ishlab chiqish;",
        "BI-RADS kategoriyasini ordinal regressiya orqali tasniflash algoritmini taklif etish;",
        "radiomik belgilarni chuqur o'rganish tasvirlovchilari bilan birlashtirish usulini qo'llash;",
        "annotatsiya, AI yordami va GPU'da model o'qitishni qo'llab-quvvatlovchi web-dasturiy majmuani yaratish;",
        "topilmalardan avtomatik klinik xulosa (hisobot) shakllantirish va uni DICOM SR sifatida eksport qilish moduli ni ishlab chiqish;",
        "dasturiy majmuani tajribaviy baholash va amaliyotda qo'llash.",
    ])

    lead(doc, "Tadqiqotning obyekti —",
         "rentgen nurlari asosida turli proyeksiyalarda (CC, MLO) olingan raqamli "
         "mammografiya tasvirlari (DICOM) va ularga tegishli klinik hisobotlar.")
    lead(doc, "Tadqiqotning predmeti —",
         "mammografiya tasvirlarida o'sma sohalarini aniqlash, BI-RADS bo'yicha tasniflash "
         "va avtomatik xulosa shakllantirish usul va algoritmlari.")
    lead(doc, "Tadqiqotning usullari.",
         "Tadqiqotda chuqur o'rganish (konvolyutsion neyron tarmoqlar, o'zaro e'tibor "
         "mexanizmi), tartibli (ordinal) regressiya, radiomik tahlil, raqamli tasvirga "
         "ishlov berish, dasturiy injiniring va ma'lumotlarni intellektual tahlil "
         "usullaridan foydalanilgan.")

    para(doc, "**Tadqiqotning ilmiy yangiligi quyidagilardan iborat:**", first_line=False)
    bullets(doc, [
        "har bir xususiyat miqyosida chap va ufqiy akslangan o'ng kontralateral xususiyatlar o'rtasida o'zaro e'tiborni hisoblovchi, o'rganiluvchi α koeffitsiyenti bilan boshqariladigan ikki tomonlama o'zaro e'tibor (BCA) bloki taklif etilgan;",
        "bir tomonning CC va MLO ko'rinishlari o'rtasida obyektlik muvofiqligini ta'minlovchi ko'rinishlararo muvofiqlik (IVC) moduli ishlab chiqilgan;",
        "BI-RADS kategoriyasini qat'iy monoton ostonalarga ega kumulyativ bog'lanish modeli orqali tasniflovchi ordinal regressiya boshi taklif etilgan;",
        "radiomik belgilarni chuqur tasvirlovchilar bilan birlashtirib, izohlanuvchanlikni oshiruvchi gibrid yondashuv qo'llanilgan;",
        "topilmalardan grounded (faqat aniqlangan ma'lumotga asoslangan) avtomatik klinik xulosa shakllantiruvchi, lokal til modelida ishlovchi modul ishlab chiqilgan.",
    ])

    para(doc, "**Tadqiqotning amaliy natijalari:**", first_line=False)
    bullets(doc, [
        "MAMOGRAF — DICOM ko'ruvchi, annotatsiya, AI inference, model o'qitish, avtomatik hisobot va PACS integratsiyasini birlashtirgan to'liq web-dasturiy majmua;",
        "GPU'da model o'qitish interfeysi (Model Studio) — jonli metrikalar va resurs monitoringi bilan;",
        "kalit talab qilmaydigan, butunlay lokal ishlovchi avtomatik hisobot generatori va DICOM SR eksporti.",
    ])

    lead(doc, "Tadqiqot natijalarining ishonchliligi",
         "algoritmlarni ishlab chiqishda chuqur o'rganish va radiomik tahlilning "
         "asoslangan matematik apparatidan foydalanilganligi, modellarning ochiq "
         "VinDr-Mammo ma'lumotlar to'plamida baholanganligi va olingan natijalarning "
         "umume'tirof etilgan metrikalar (mAP, F1, IoU) bilan tasdiqlanganligi bilan "
         "izohlanadi.")
    lead(doc, "Tadqiqotning ilmiy va amaliy ahamiyati.",
         "Ilmiy ahamiyati — mammografiyada ikki tomonlama asimmetriya, ko'pko'rinishli "
         "muvofiqlik va ordinal baholash kabi klinik induktiv taxminlarni rasmiy "
         "matematik shaklda ifodalashdan iborat. Amaliy ahamiyati — ishlab chiqilgan "
         "MAMOGRAF dasturiy majmuasini radiologiya bo'limlarida tashxis jarayonini "
         "tezlashtirish va xatolarni kamaytirish uchun qo'llash imkoniyatidir.")
    lead(doc, "Tadqiqot natijalarining joriy qilinishi.",
         "[Dasturiy majmua ... klinikasi/poliklinikasi radiologiya bo'limida joriy "
         "qilingan — joriy etilish ma'lumotlarini to'ldiring.]")
    lead(doc, "Dissertatsiyaning tuzilishi va hajmi.",
         "Dissertatsiya kirish, to'rtta bob, xulosa, foydalanilgan adabiyotlar ro'yxati "
         "va ilovalardan iborat.")
    doc.add_page_break()

    # ===== I BOB =====
    h1(doc, "I bob. Mammografiya tasvirlarini tahlil qilishning zamonaviy holati")

    h2(doc, "1.1-§. Koʻkrak saratoni, skrining mammografiyasi va BI-RADS tizimi")
    para(doc, "Ko'krak saratoni — ko'krak bezi to'qimasida xavfli o'smaning rivojlanishi "
              "bilan tavsiflanuvchi kasallik. Uni erta aniqlashning asosiy vositasi — "
              "skrining mammografiyasi bo'lib, protokol bo'yicha har bir tadqiqot uchun "
              "to'rtta standart ko'rinish hosil qilinadi: har bir ko'krakning kraniokaudal "
              "(CC) va mediolateral qiyshiq (MLO) proyeksiyalari. Radiolog bu ko'rinishlarni "
              "birgalikda baholaydi: har bir proyeksiyada chap va o'ng ko'krakni solishtirib "
              "asimmetriyani aniqlaydi, hamda bir ko'krakning CC va MLO ko'rinishlarini "
              "birlashtirib shikastlanishni fazoda joylashtiradi.")
    para(doc, "Topilmalarni standartlashtirish uchun Amerika Radiologiya Kolleji (ACR) "
              "tomonidan ishlab chiqilgan **BI-RADS** (Breast Imaging Reporting and Data "
              "System) tizimi qo'llaniladi. U topilmaning xavflilik darajasini 0 dan 6 "
              "gacha bo'lgan kategoriyalar bilan ifodalaydi (1.1-jadval). Muhim jihat — "
              "BI-RADS shkalasi **tartibli (ordinal)**: kategoriyalar orasidagi masofa "
              "teng emas va xatolarning klinik narxi turlicha.")
    table(doc, ["BI-RADS", "Talqin", "Tavsiya"], [
        ["0", "Baho to'liq emas", "Qo'shimcha tasvirlash"],
        ["1", "Patologiya yo'q", "Rutin skrining"],
        ["2", "Benign (xavfsiz)", "Rutin skrining"],
        ["3", "Ehtimol benign", "6 oydan keyin nazorat"],
        ["4 (A/B/C)", "Shubhali", "Biopsiya"],
        ["5", "Xavfli ehtimoli yuqori", "Biopsiya / konsultatsiya"],
        ["6", "Tasdiqlangan rak", "Davolash"],
    ], caption="1.1-jadval. BI-RADS kategoriyalari va klinik tavsiyalar")

    h2(doc, "1.2-§. Tibbiy tasvirlarga ishlov berish va DICOM standarti")
    para(doc, "Raqamli tibbiy tasvirlar **DICOM** (Digital Imaging and Communications in "
              "Medicine) standartida saqlanadi. DICOM fayli piksel ma'lumotlari bilan "
              "birga bemor, tadqiqot, qurilma va proyeksiya (ViewPosition, ImageLaterality) "
              "haqidagi metama'lumotlarni o'z ichiga oladi. Klinik tizimlar (PACS) "
              "tasvirlarni saqlash va almashish uchun DICOM tarmoq xizmatlaridan "
              "(C-ECHO, C-FIND, C-STORE, C-MOVE) foydalanadi. Tasvirlar tahlilga "
              "tayyorlanishidan oldin shaxsiy ma'lumotlar (PHI) anonimlashtirilishi — "
              "maxfiylik talablariga muvofiq — zarur.")
    para(doc, "Tasvir sifatini oshirish bosqichida normalizatsiya, kontrast moslash "
              "(masalan, CLAHE), shovqinni kamaytirish va o'lchamni moslash kabi dastlabki "
              "ishlov berish usullari qo'llaniladi.")

    h2(doc, "1.3-§. Chuqur oʻrganish va radiomika asosida oʻsma aniqlash usullari")
    para(doc, "Obyektlarni aniqlashda ikki asosiy yo'nalish mavjud: ikki bosqichli "
              "(Faster R-CNN) va bir bosqichli (YOLO, SSD) detektorlar. YOLO oilasi "
              "yuqori tezligi va aniqligi bilan tibbiy tasvirlashda keng qo'llaniladi. "
              "Aniqlik metrikalari sifatida IoU, precision, recall, F1 va mAP ishlatiladi "
              "(2-bobda batafsil). O'zaro e'tibor (cross-attention) mexanizmi ikki "
              "xususiyat oqimini bog'lash uchun samarali vosita bo'lib, ko'p modalli "
              "birlashtirishda qo'llaniladi.")
    para(doc, "**Radiomika** — tibbiy tasvirdan ko'p sonli miqdoriy belgilarni (shakl, "
              "intensivlik, tekstura) avtomatik chiqarib, ularni tahlil qilish yo'nalishi. "
              "Radiomik belgilar chuqur o'rganish modellariga izohlanuvchanlik (explainability) "
              "qo'shadi va \"qora quti\" muammosini yumshatadi.")

    h2(doc, "1.4-§. Tadqiqot masalasining qoʻyilishi")
    para(doc, "Yuqorida keltirilgan tahlil asosida quyidagi ilmiy masala qo'yiladi: "
              "mammografiyaning to'rtta ko'rinishini birgalikda, ikki tomonlama asimmetriya "
              "va ko'rinishlararo muvofiqlikni hisobga olgan holda qayta ishlovchi, "
              "BI-RADS ni ordinal shkalada baholovchi va topilmalardan avtomatik klinik "
              "xulosa shakllantiruvchi yagona dasturiy majmuani ishlab chiqish.")
    chapter_concl(doc, "I", [
        "Ko'krak saratonini erta aniqlashda skrining mammografiyasi va BI-RADS tizimi asosiy o'rin tutadi; BI-RADS ordinal tabiatga ega.",
        "Mavjud chuqur o'rganish detektorlari to'rt ko'rinishni mustaqil qayta ishlab, ikki tomonlama va ko'rinishlararo signallardan to'liq foydalanmaydi.",
        "Radiomik tahlil chuqur o'rganishga izohlanuvchanlik qo'shadi; ikki yondashuvni birlashtirish istiqbolli.",
        "Aniqlangan kamchiliklar BCA-YOLO arxitekturasi va MAMOGRAF dasturiy majmuasini ishlab chiqishni asoslab beradi.",
    ])
    doc.add_page_break()

    # ===== II BOB =====
    h1(doc, "II bob. Koʻkrak oʻsmalarini aniqlash va tasniflash algoritmlari")

    h2(doc, "2.1-§. BCA-YOLO arxitekturasi: ikki tomonlama oʻzaro eʼtibor va koʻrinishlararo muvofiqlik")
    para(doc, "Ushbu ishda **BCA-YOLO** arxitekturasi taklif etiladi — bu YOLO oilasiga "
              "uchta yangi komponent kiritilgan model: (i) ikki tomonlama o'zaro e'tibor "
              "(BCA) bloki, (ii) ko'rinishlararo muvofiqlik (IVC) moduli, (iii) ordinal "
              "BI-RADS regressiya boshi. Umumiy arxitektura 2.1-rasmda keltirilgan.")
    img(doc, "arch_bca_yolo.png", width=5.6, caption="2.1-rasm. BCA-YOLO arxitekturasi")
    para(doc, "**Ikki tomonlama o'zaro e'tibor (BCA).** Har bir xususiyat miqyosi $s$ da "
              "chap ko'rinish xususiyati $F_L$ va ufqiy akslangan o'ng ko'rinish xususiyati "
              "$F_R$ o'rtasida o'zaro e'tibor hisoblanadi. So'rov $Q$ chapdan, kalit $K$ va "
              "qiymat $V$ o'ngdan olinadi:")
    eq(doc, "bca_attn", number="2.1")
    para(doc, "Natija o'rganiluvchi skalar $\\alpha$ koeffitsiyenti bilan qoldiq (residual) "
              "tarzda qo'shiladi. $\\alpha$ nolga initsializatsiya qilinadi — bu o'rganish "
              "boshida modelni aynan bazaviy YOLO ga teng qiladi va oldindan o'rgatilgan "
              "tarmoqlar bilan mosligini saqlaydi:")
    eq(doc, "bca_res", number="2.2")
    para(doc, "Asimmetriya xaritasi $\\bar{A}^{(s)}$ shikastlanish joylarida cho'qqiga "
              "ko'tarilishga o'rgatiladi:")
    eq(doc, "asym", number="2.3")
    para(doc, "**Ko'rinishlararo muvofiqlik (IVC).** Bir tomonning CC va MLO ko'rinishlari "
              "uchun obyektlik logitining global o'rtachasi olinadi va ular o'rtasida "
              "silliqlik (muvofiqlik) talab qilinadi — bir shikastlanish ikkala "
              "proyeksiyadan ham ko'rinishi kerak:")
    eq(doc, "ivc_gap", number="2.4")
    eq(doc, "ivc", number="2.5")

    h2(doc, "2.2-§. Ordinal BI-RADS regressiyasi va kompozit yoʻqotish funksiyasi")
    para(doc, "BI-RADS shkalasining ordinal tabiatini hisobga olish uchun nominal "
              "kross-entropiya o'rniga **kumulyativ bog'lanish** modeli qo'llaniladi. "
              "Maqsad $y \\in \\{0,\\dots,K-1\\}$ uchun ordinal yo'qotish ikkilik "
              "kross-entropiyalar yig'indisi sifatida ifodalanadi:")
    eq(doc, "ordinal", number="2.6")
    para(doc, "Ostonalar $\\theta_k$ ketma-ketligi $\\mathrm{softplus}$ orqali qat'iy "
              "monoton qilib parametrlanadi — bu qo'shimcha jarimasiz monotonlikni "
              "kafolatlaydi va qo'shni BI-RADS darajalarining yagona chegaraga "
              "yig'ilishini oldini oladi:")
    eq(doc, "thresh", number="2.7")
    para(doc, "BCA darvozasi $\\alpha$ bo'yicha gradiyent initsializatsiya paytida ham "
              "nol bo'lmaydi (Frobenius ichki ko'paytmasi orqali), shu sababli optimizator "
              "barqaror tushish yo'nalishini oladi:")
    eq(doc, "grad_alpha", number="2.8")
    para(doc, "Umumiy (kompozit) yo'qotish funksiyasi YOLO ning standart komponentlari "
              "(quti, sinf, taqsimot fokal yo'qotishi) bilan taklif etilgan uch "
              "regulyarizatsiyani birlashtiradi:")
    eq(doc, "total", number="2.9")
    para(doc, "Namunaviy realizatsiya YOLOv8/v11 bazaviy tarmog'iga taxminan **3,1 mln** "
              "o'rganiluvchi parametr qo'shadi (bazaviy modelning ~5–8 foizi). BCA "
              "hisoblash miqyoslar va boshlar bo'yicha parallellashtirilgani uchun "
              "inferens vaqtiga deyarli ta'sir qilmaydi.")

    h2(doc, "2.3-§. Radiomik belgilarni chuqur oʻrganish bilan birlashtirish")
    para(doc, "Izohlanuvchanlikni oshirish uchun chuqur tasvirlovchilar radiomik belgilar "
              "bilan birlashtiriladi (gibrid yondashuv). Radiomik belgilar orasida "
              "tekstura tahlili uchun kulrang darajalar birgalikdagi paydo bo'lish "
              "matritsasidan (GLCM) olinadigan kontrast va entropiya muhim o'rin tutadi:")
    eq(doc, "glcm", number="2.10")
    para(doc, "Chuqur xususiyat vektori $f_{deep}$ va radiomik vektor $f_{radiomic}$ "
              "birlashtirilib (konkatenatsiya), to'liq bog'langan qatlam orqali tasniflanadi:")
    eq(doc, "fusion", number="2.11")
    para(doc, "Bu gibrid yondashuv chuqur o'rganishning yuqori aniqligini radiomikaning "
              "patologik asoslangan izohlanuvchanligi bilan birlashtiradi.")
    chapter_concl(doc, "II", [
        "Ikki tomonlama o'zaro e'tibor (BCA) bloki kontralateral asimmetriyani o'rganiluvchi α koeffitsiyenti bilan modelga kiritadi.",
        "Ko'rinishlararo muvofiqlik (IVC) moduli CC va MLO ko'rinishlari o'rtasida obyektlik muvofiqligini ta'minlaydi.",
        "Ordinal regressiya boshi BI-RADS ni monoton ostonalar bilan kumulyativ tarzda baholaydi.",
        "Radiomik belgilarni chuqur tasvirlovchilar bilan birlashtirish izohlanuvchanlikni oshiradi.",
    ])
    doc.add_page_break()

    # ===== III BOB =====
    h1(doc, "III bob. MAMOGRAF dasturiy majmuasini amalga oshirish")

    h2(doc, "3.1-§. Tizim arxitekturasi, maxfiylik va xavfsizlik")
    para(doc, "MAMOGRAF — taklif etilgan algoritmlarni amaliyotga tatbiq etuvchi web-dasturiy "
              "majmua. Backend FastAPI (Python) da, frontend HTML/JavaScript da amalga "
              "oshirilgan; DICOM bilan ishlash pydicom orqali, AI Ultralytics YOLO va "
              "PyTorch (CUDA) orqali bajariladi. Umumiy arxitektura 3.1-rasmda.")
    img(doc, "architecture.png", width=6.0, caption="3.1-rasm. MAMOGRAF tizim arxitekturasi")
    para(doc, "**Maxfiylik:** har bir yuklangan DICOM faylda shaxsiy ma'lumotlar (PHI — "
              "ism, ID, sana, shifokor) avtomatik anonimlashtiriladi; keyingi barcha "
              "bosqichlar PHI-siz nusxa ustida bajariladi. **Xavfsizlik:** JWT "
              "autentifikatsiya, rol asosidagi ruxsatlar (admin/reviewer/annotator) va "
              "ixtiyoriy ikki bosqichli (TOTP) autentifikatsiya qo'llaniladi.")

    h2(doc, "3.2-§. Annotatsiya, AI yordami va model oʻqitish (Model Studio)")
    para(doc, "Radiolog tasvirda shubhali sohalarni quti yoki ko'pburchak bilan belgilaydi, "
              "har biriga sinf, BI-RADS kategoriyasi va izoh biriktiradi (3.2-rasm). "
              "Annotatsiyalar holat oqimi (draft → submitted → approved/rejected) orqali "
              "tasdiqlanadi. AI yordami: avto-inference, smart-click va noaniqlik "
              "(uncertainty) issiqlik xaritasi.")
    img(doc, "ui_viewer.png", width=6.2, caption="3.2-rasm. DICOM ko'ruvchi va annotatsiya")
    para(doc, "**Model Studio** — o'z annotatsiyalaringizdan yangi YOLO modelini GPU'da "
              "o'qitish interfeysi (3.3-rasm). U dataset tekshiruvi, to'liq giperparametrlar, "
              "jonli grafiklar (loss, mAP, precision/recall), GPU monitoringi (VRAM, "
              "utilization, harorat) va o'qitishni boshqarish (Train/Stop/Resume) "
              "imkoniyatlarini beradi. O'qitilgan model avtomatik joylashtiriladi.")
    img(doc, "ui_studio.png", width=6.2, caption="3.3-rasm. Model Studio — GPU'da o'qitish")

    h2(doc, "3.3-§. Avtomatik xulosa shakllantirish va DICOM SR/PACS integratsiyasi")
    para(doc, "Belgilangan lezyonlardan strukturaviy topilmalar (laterallik, kvadrant, "
              "tur, o'lcham, BI-RADS) shakllantirilib, ulardan avtomatik klinik xulosa "
              "(hisobot) yaratiladi (3.4-rasm). Xulosa ikki usulda hosil bo'ladi: "
              "deterministik shablon yoki **lokal til modeli** (Ollama, qwen2.5) — bu "
              "tashqi API kaliti va internetni talab qilmaydi. Model **faqat berilgan "
              "topilmalardan** foydalanadi (ma'lumot to'qimaydi). Chiqish doim radiolog "
              "tasdiqlaydigan qoralama bo'lib, **DICOM SR** (Structured Report) sifatida "
              "eksport qilinadi va PACS'ga jo'natiladi.")
    img(doc, "ui_report.png", width=6.2, caption="3.4-rasm. Avtomatik hisobot generatori (lokal AI)")
    chapter_concl(doc, "III", [
        "MAMOGRAF FastAPI backend va web-interfeysda amalga oshirilgan; PHI avtomatik anonimlashtiriladi.",
        "Annotatsiya, AI yordami va GPU'da model o'qitish (Model Studio) yagona tizimda birlashtirilgan.",
        "Avtomatik xulosa lokal til modelida grounded tarzda shakllantiriladi va DICOM SR sifatida eksport qilinadi.",
        "Tizim PACS bilan integratsiyalashgan va rol asosidagi xavfsizlikka ega.",
    ])
    doc.add_page_break()

    # ===== IV BOB =====
    h1(doc, "IV bob. Tajribaviy tadqiqotlar va amaliyotda qoʻllash")

    h2(doc, "4.1-§. Maʼlumotlar toʻplami va baholash metodologiyasi")
    para(doc, "Modellarni baholash uchun ochiq **VinDr-Mammo** ma'lumotlar to'plami "
              "(taxminan 20 000 tasvir) ishlatilgan. Aniqlik metrikalari sifatida quyidagilar "
              "qo'llaniladi. Aniqlangan va haqiqiy quti o'rtasidagi mos kelish IoU bilan:")
    eq(doc, "iou", number="4.1")
    para(doc, "Precision, recall va F1 ko'rsatkichlari:")
    eq(doc, "prf", number="4.2")
    para(doc, "Detektsiya sifatining umumiy ko'rsatkichi — o'rtacha aniqlik (mAP):")
    eq(doc, "map", number="4.3")
    para(doc, "Segmentatsiya sifati uchun Dice koeffitsiyenti:")
    eq(doc, "dice", number="4.4")

    h2(doc, "4.2-§. Tajriba natijalari va tahlili")
    para(doc, "Radiomik belgilarni chuqur o'rganish bilan birlashtiruvchi gibrid yondashuv "
              "VinDr-Mammo to'plamida quyidagi natijalarni ko'rsatdi (4.1-jadval).")
    table(doc, ["Ko'rsatkich", "Qiymat"], [
        ["Aniqlik (Accuracy)", "0,904"],
        ["F1-ball", "0,89"],
        ["mAP", "0,78"],
    ], caption="4.1-jadval. Gibrid (radiomika + chuqur o'rganish) model natijalari (VinDr-Mammo)")
    para(doc, "BCA-YOLO arxitekturasining namunaviy realizatsiyasi bazaviy YOLO ga "
              "taxminan 3,1 mln parametr qo'shadi va $\\alpha=0$ da bazaviy modelni aynan "
              "tiklaydi (sog'liq tekshiruvi bilan tasdiqlangan). To'liq ablatsion tahlil "
              "(BCA, IVC, ordinal komponentlarining alohida hissasi) va radiologlar paneli "
              "bilan klinik baholash davom etmoqda.")
    para(doc, "**[Eslatma: quyidagi jadvalni o'zingizning to'liq tajriba natijalaringiz "
              "bilan to'ldiring — ablatsion tahlil, har komponentning mAP/F1 ga hissasi, "
              "bazaviy modellar bilan taqqoslash.]**", first_line=False)
    table(doc, ["Konfiguratsiya", "Precision", "Recall", "F1", "mAP@50"], [
        ["Bazaviy YOLO", "[…]", "[…]", "[…]", "[…]"],
        ["+ BCA", "[…]", "[…]", "[…]", "[…]"],
        ["+ BCA + IVC", "[…]", "[…]", "[…]", "[…]"],
        ["+ BCA + IVC + Ordinal (to'liq)", "[…]", "[…]", "[…]", "[…]"],
    ], caption="4.2-jadval. Ablatsion tahlil (to'ldiriladi)")

    h2(doc, "4.3-§. Dasturiy majmuani amaliyotda qoʻllash")
    para(doc, "Ishlab chiqilgan MAMOGRAF dasturiy majmuasi radiologiya bo'limida tashxis "
              "jarayonini qo'llab-quvvatlash uchun mo'ljallangan: tasvir yuklash → AI yordami "
              "yoki qo'lda annotatsiya → reviewer tasdig'i → avtomatik hisobot → DICOM SR "
              "eksporti. Tizim PACS bilan integratsiyalashgani uchun mavjud klinik "
              "infratuzilmaga moslashadi. Statistika va audit modullari sifat nazorati va "
              "kuzatuvchanlikni ta'minlaydi.")
    chapter_concl(doc, "IV", [
        "Modellar ochiq VinDr-Mammo to'plamida umume'tirof etilgan metrikalar bilan baholandi.",
        "Gibrid yondashuv aniqlik 0,904, F1 0,89 va mAP 0,78 natijalarini ko'rsatdi.",
        "BCA-YOLO bazaviy modelga ~3,1 mln parametr qo'shib, inferens vaqtiga deyarli ta'sir qilmaydi.",
        "MAMOGRAF dasturiy majmuasi amaliyotda to'liq ish oqimini qo'llab-quvvatlaydi.",
    ])
    doc.add_page_break()

    # ===== XULOSA =====
    h1(doc, "Xulosa")
    para(doc, "Dissertatsiya ishi natijasida mammografiya tasvirlarida ko'krak o'sma "
              "sohalarini aniqlash, BI-RADS bo'yicha tasniflash va avtomatik klinik xulosa "
              "shakllantirishning sun'iy intellektga asoslangan dasturiy majmuasi ishlab "
              "chiqildi. Asosiy natijalar:")
    bullets(doc, [
        "Mammografiyaning klinik induktiv taxminlarini — ikki tomonlama asimmetriya, ko'rinishlararo muvofiqlik va ordinal BI-RADS baholashni — rasmiy matematik shaklda ifodalovchi BCA-YOLO arxitekturasi taklif etildi.",
        "BCA bloki o'rganiluvchi α koeffitsiyenti bilan kontralateral asimmetriyani modelga kiritadi va o'rganish boshida bazaviy modelni aynan tiklaydi (barqaror sozlash).",
        "Ordinal regressiya boshi BI-RADS ni qat'iy monoton ostonalar bilan kumulyativ tarzda baholaydi.",
        "Radiomik belgilarni chuqur o'rganish bilan birlashtirish izohlanuvchanlikni oshirdi (VinDr-Mammo'da aniqlik 0,904, F1 0,89, mAP 0,78).",
        "Annotatsiya, AI yordami, GPU'da model o'qitish (Model Studio), avtomatik hisobot va PACS integratsiyasini birlashtirgan MAMOGRAF dasturiy majmuasi yaratildi.",
        "Avtomatik hisobot generatori butunlay lokal (Ollama) ishlaydi, tashqi xizmatga bog'liq emas va grounded — ma'lumot to'qimaydi; natija DICOM SR sifatida eksport qilinadi.",
    ], numbered=True)
    para(doc, "Tadqiqot natijalari mammografik tashxis jarayonini tezlashtirish va "
              "xatolarni kamaytirishga, shu orqali ko'krak saratonini erta aniqlashga "
              "hissa qo'shadi.")
    doc.add_page_break()

    # ===== ADABIYOTLAR =====
    h1(doc, "Foydalanilgan adabiyotlar roʻyxati")
    refs = [
        "World Health Organization. Breast cancer. — Geneva: WHO, 2023.",
        "Sung H. et al. Global Cancer Statistics 2020: GLOBOCAN // CA: A Cancer Journal for Clinicians. — 2021. — Vol. 71. — P. 209–249.",
        "American College of Radiology. Breast Imaging Reporting and Data System (BI-RADS). — 5th ed. — ACR, 2013.",
        "Redmon J., Farhadi A. YOLOv3: An Incremental Improvement // arXiv:1804.02767. — 2018.",
        "Jocher G. et al. Ultralytics YOLO (v8/v11). — 2023–2024.",
        "Ren S., He K., Girshick R., Sun J. Faster R-CNN: Towards Real-Time Object Detection // NeurIPS. — 2015.",
        "Vaswani A. et al. Attention Is All You Need // NeurIPS. — 2017.",
        "He K., Zhang X., Ren S., Sun J. Deep Residual Learning for Image Recognition // CVPR. — 2016.",
        "Lin T.-Y. et al. Focal Loss for Dense Object Detection (RetinaNet) // ICCV. — 2017.",
        "Tian Z. et al. FCOS: Fully Convolutional One-Stage Object Detection // ICCV. — 2019.",
        "Lambin P. et al. Radiomics: extracting more information from medical images // European Journal of Cancer. — 2012. — Vol. 48. — P. 441–446.",
        "van Griethuysen J.J.M. et al. Computational Radiomics System (PyRadiomics) // Cancer Research. — 2017. — Vol. 77. — P. 104–107.",
        "Nguyen H.T. et al. VinDr-Mammo: A large-scale benchmark dataset for computer-aided diagnosis in full-field digital mammography // Scientific Data. — 2023.",
        "Haralick R.M. et al. Textural Features for Image Classification // IEEE Trans. on Systems, Man, and Cybernetics. — 1973.",
        "Niu Z. et al. Ordinal Regression with Multiple Output CNN for Age Estimation // CVPR. — 2016.",
        "Mason D. et al. pydicom: An open source DICOM library. — 2008–2024.",
        "Ramaswamy S. et al. FastAPI: Modern, fast web framework for building APIs. — 2018–2024.",
        "Paszke A. et al. PyTorch: An Imperative Style, High-Performance Deep Learning Library // NeurIPS. — 2019.",
        "[Muallif] BCA-YOLO: Bilateral Cross-Attention with Inter-View Consistency and Ordinal BI-RADS Regression for Mammographic Lesion Detection. — 2024.",
        "[Muallif] Radiomic-Enhanced Faster R-CNN Framework for Breast Cancer Detection in Mammography. — 2024.",
        "O'zbekiston Respublikasi Prezidentining «Raqamli O'zbekiston-2030» strategiyasini tasdiqlash to'g'risidagi Farmoni. — 2020.",
    ]
    for i, r in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.left_indent = Inches(0.3)
        p.add_run(f"{i}. ").bold = True
        p.add_run(r).font.size = Pt(12)
    para(doc, "[Eslatma: adabiyotlar ro'yxatini o'z manbalaringiz va GOST/talab etilgan "
              "uslubga ko'ra to'ldiring/tahrirlang.]", italic=True, muted=True, first_line=False)
    doc.add_page_break()

    # ===== ILOVALAR =====
    h1(doc, "Ilovalar")
    img(doc, "ui_login.png", width=6.0, caption="A ilova. Tizimga kirish")
    img(doc, "ui_datasetprep.png", width=6.0, caption="B ilova. Dataset tayyorlash")
    img(doc, "ui_models.png", width=6.0, caption="C ilova. Modellar dashboard")
    img(doc, "ui_stats.png", width=6.0, caption="D ilova. Statistika")
    img(doc, "ui_pacs.png", width=6.0, caption="E ilova. PACS integratsiyasi")
    img(doc, "arch_tillnet.png", width=5.6, caption="F ilova. TILLNet-Det (rasm+matn) arxitekturasi")

    doc.save(str(OUT))
    print("Saqlandi:", OUT)


if __name__ == "__main__":
    build()
