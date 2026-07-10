# -*- coding: utf-8 -*-
"""maqola_kit.py — ilmiy maqola .docx qurish uchun umumiy yordamchilar.

Ikkala 2026-yilgi maqola (barqarorlik + ansambl) shu modulni ishlatadi.
Matn ichidagi `$...$` bo'laklari ham Word-native OMML tenglamaga aylantiriladi
(omml_inline.emit_rich orqali), `**...**` — qalin.
"""
from __future__ import annotations

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt, RGBColor

import omml_inline as _inline

ACCENT = RGBColor(0x1A, 0x3E, 0x6E)
GREY = RGBColor(0x33, 0x33, 0x33)


# ───────────────────────────── formatlash ──────────────────────────────── #

def N(x, lang, nd=3):
    """Sonni tilga mos o'nlik ajratgich bilan."""
    if x is None or (isinstance(x, float) and x != x):
        return "—"
    s = f"{x:.{nd}f}"
    return s.replace(".", ",") if lang in ("uz", "ru") else s


def CI(v, ci, lang, nd=3):
    if not ci:
        return N(v, lang, nd)
    return f"{N(v, lang, nd)} [{N(ci['lo'], lang, nd)}; {N(ci['hi'], lang, nd)}]"


def PV(p, lang):
    if p is None:
        return "—"
    if p < 0.001:
        return "< 0,001" if lang in ("uz", "ru") else "< 0.001"
    return N(p, lang, 3)


def PP(x, lang, nd=1):
    """Foizli punkt (percentage point) ko'rinishida, ishorasi bilan."""
    s = f"{x * 100:+.{nd}f}"
    return s.replace(".", ",") if lang in ("uz", "ru") else s


def INT(n, lang):
    """Ming ajratgichli butun son."""
    s = f"{int(n):,}".replace(",", " ")
    return s


# ─────────────────────────── hujjat elementlari ────────────────────────── #

def setup(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(0)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2.0)


def para(doc, text, first_indent=True, size=12, align=None, italic=False):
    p = doc.add_paragraph()
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_indent:
        p.paragraph_format.first_line_indent = Cm(1.25)
    _inline.emit_rich(p, text, size=size, italic=italic)
    return p


def title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = ACCENT
    return p


def authors(doc, lines):
    for i, ln in enumerate(lines):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2 if i < len(lines) - 1 else 10)
        r = p.add_run(ln)
        r.font.size = Pt(11 if i == 0 else 10)
        r.bold = (i == 0)
        r.italic = (i > 0)


def heading(doc, text, size=13):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    _inline.emit_rich(p, text, size=size, bold_all=True, color=ACCENT)
    return p


def subheading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    _inline.emit_rich(p, text, size=11.5, bold_all=True, italic=True)
    return p


def abstract(doc, label, text, kw_label, kw):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(label + " ")
    r.bold = True
    r.font.size = Pt(10.5)
    _inline.emit_rich(p, text, size=10.5)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p2.paragraph_format.space_after = Pt(10)
    r = p2.add_run(kw_label + " ")
    r.bold = True
    r.font.size = Pt(10.5)
    r = p2.add_run(kw)
    r.italic = True
    r.font.size = Pt(10.5)


def caption(doc, text, before=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6 if before else 2)
    p.paragraph_format.space_after = Pt(8 if not before else 2)
    _inline.emit_rich(p, text, size=9.5, italic=True, color=GREY)


def table(doc, header, rows, size=8.5, first_left=True):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _inline.emit_rich(c.paragraphs[0], str(h), size=size, bold_all=True)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            cells[i].paragraphs[0].alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if (i == 0 and first_left)
                else WD_ALIGN_PARAGRAPH.CENTER)
            _inline.emit_rich(cells[i].paragraphs[0], str(v), size=size)
    for r_ in t.rows:
        for c_ in r_.cells:
            c_.paragraphs[0].paragraph_format.space_after = Pt(0)
            c_.paragraphs[0].paragraph_format.line_spacing = 1.0
    return t


def picture(doc, path, cap, width=6.1):
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.space_before = Pt(6)
    caption(doc, cap)


def refs(doc, label, items):
    heading(doc, label)
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(f"{i}. {it}")
        r.font.size = Pt(10)
