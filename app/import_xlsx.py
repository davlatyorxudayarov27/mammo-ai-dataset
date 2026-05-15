from __future__ import annotations

import argparse
import sys
import time
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook

from .db import DB_PATH, get_conn, init_db

XLSX_TO_DB = {
    "hasta_adi": "first_name",
    "soyadi": "last_name",
    "kimlik_no": "patient_id",
    "jins": "sex",
    "tugilgan": "birth_date",
    "hizmettarihi": "service_date",
    "kod_ara": "exam_code",
    "tetkik_ismi": "exam_name",
    "onay_tarihi": "approval_date",
    "Mamologiya Report": "report",
    "hikayesi": "history",
    "sikayeti": "complaints",
    "klinik_bulgular": "clinical_findings",
    "oneriler": "recommendations",
    "hastalikhakkindabilgi": "disease_info",
    "tedavi": "treatment",
    "laboratuvar_notu": "lab_notes",
    "taburcu_ilaclari": "discharge_meds",
    "mesleki_anemnez": "occupational",
    "gelis_nedeni": "admission_reason",
    "tani": "diagnosis",
    "diagnostik_uygulamalar": "diagnostic_procedures",
}

RECORD_FIELDS = [
    "source_row", "patient_id",
    "service_date", "approval_date",
    "exam_code", "exam_name",
    "report", "history", "complaints", "clinical_findings",
    "recommendations", "disease_info", "treatment", "lab_notes",
    "discharge_meds", "occupational", "admission_reason",
    "diagnosis", "diagnostic_procedures",
    "imported_at",
]

EXCEL_EPOCH = datetime(1899, 12, 30)


def to_text(v):
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return s or None
    return str(v).strip() or None


def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    try:
        s = float(v)
    except (TypeError, ValueError):
        return None
    if s <= 0 or s > 80000:
        return None
    try:
        return (EXCEL_EPOCH + timedelta(days=s)).date().isoformat()
    except OverflowError:
        return None


def import_xlsx(path: Path, reset: bool = False) -> dict:
    init_db()
    if reset:
        with get_conn() as c:
            c.execute("DELETE FROM records")
            c.execute("DELETE FROM patients")
            c.commit()

    t0 = time.perf_counter()
    wb = load_workbook(str(path), read_only=True, data_only=True)
    ws = wb.active

    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if header is None:
        wb.close()
        raise ValueError("xlsx is empty")
    col_idx = {name: i for i, name in enumerate(header) if name}

    missing_cols = [k for k in XLSX_TO_DB if k not in col_idx]
    if missing_cols:
        print(f"[warn] missing columns in xlsx: {missing_cols}", file=sys.stderr)

    def cell(row, col_name):
        i = col_idx.get(col_name)
        if i is None or i >= len(row):
            return None
        return row[i]

    patients: dict[str, tuple] = {}
    records: list[dict] = []
    skipped_empty = 0
    now = datetime.now(timezone.utc).isoformat()

    for r_idx, row in enumerate(rows_iter, start=2):
        patient_id = to_text(cell(row, "kimlik_no"))
        if not patient_id:
            skipped_empty += 1
            continue

        if patient_id not in patients:
            patients[patient_id] = (
                to_text(cell(row, "hasta_adi")),
                to_text(cell(row, "soyadi")),
                to_text(cell(row, "jins")),
                to_date(cell(row, "tugilgan")),
            )

        rec = {
            "source_row": r_idx,
            "patient_id": patient_id,
            "service_date": to_date(cell(row, "hizmettarihi")),
            "approval_date": to_date(cell(row, "onay_tarihi")),
            "exam_code": to_text(cell(row, "kod_ara")),
            "exam_name": to_text(cell(row, "tetkik_ismi")),
            "report": to_text(cell(row, "Mamologiya Report")),
            "history": to_text(cell(row, "hikayesi")),
            "complaints": to_text(cell(row, "sikayeti")),
            "clinical_findings": to_text(cell(row, "klinik_bulgular")),
            "recommendations": to_text(cell(row, "oneriler")),
            "disease_info": to_text(cell(row, "hastalikhakkindabilgi")),
            "treatment": to_text(cell(row, "tedavi")),
            "lab_notes": to_text(cell(row, "laboratuvar_notu")),
            "discharge_meds": to_text(cell(row, "taburcu_ilaclari")),
            "occupational": to_text(cell(row, "mesleki_anemnez")),
            "admission_reason": to_text(cell(row, "gelis_nedeni")),
            "diagnosis": to_text(cell(row, "tani")),
            "diagnostic_procedures": to_text(cell(row, "diagnostik_uygulamalar")),
            "imported_at": now,
        }
        records.append(rec)

    wb.close()

    with get_conn() as c:
        c.executemany(
            "INSERT OR REPLACE INTO patients "
            "(patient_id, first_name, last_name, sex, birth_date) VALUES (?,?,?,?,?)",
            [(pid, *vals) for pid, vals in patients.items()],
        )
        if records:
            placeholders = ",".join("?" * len(RECORD_FIELDS))
            c.executemany(
                f"INSERT OR REPLACE INTO records ({','.join(RECORD_FIELDS)}) "
                f"VALUES ({placeholders})",
                [tuple(r[k] for k in RECORD_FIELDS) for r in records],
            )
        c.commit()

    return {
        "path": str(path),
        "patients_unique": len(patients),
        "records_inserted": len(records),
        "skipped_empty_rows": skipped_empty,
        "elapsed_s": round(time.perf_counter() - t0, 2),
        "db_path": str(DB_PATH),
    }


def main():
    ap = argparse.ArgumentParser(description="Import MamologiyaInfo xlsx into SQLite.")
    default_path = Path(__file__).resolve().parent.parent / "MamologiyaInfo_.xlsx"
    ap.add_argument("path", nargs="?", default=str(default_path),
                    help="path to xlsx file (default: ../MamologiyaInfo_.xlsx)")
    ap.add_argument("--reset", action="store_true",
                    help="drop existing rows before import")
    args = ap.parse_args()
    p = Path(args.path)
    if not p.exists():
        print(f"not found: {p}", file=sys.stderr)
        sys.exit(1)
    result = import_xlsx(p, reset=args.reset)
    width = max(len(k) for k in result)
    for k, v in result.items():
        print(f"  {k.ljust(width)} : {v}")


if __name__ == "__main__":
    main()
