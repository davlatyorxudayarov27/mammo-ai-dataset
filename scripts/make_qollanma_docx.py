# -*- coding: utf-8 -*-
"""MAMOGRAF — To'liq foydalanuvchi qo'llanmasi (.docx).

QOLLANMA.md mazmunini Word hujjatiga (rasmlar + jadvallar bilan) aylantiradi.
Skrinshotlar doc_assets/ da bo'lishi kerak (make_full_screenshots.py).
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "doc_assets"
OUT = ROOT / "MAMOGRAF_qollanma.docx"

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x55, 0x55, 0x55)
WARN_BG = "FDEDEC"


def setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for lvl, sz in ((1, 17), (2, 13), (3, 12)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = "Calibri"
        st.font.size = Pt(sz)
        st.font.color.rgb = ACCENT
        st.font.bold = True


def h1(doc, t):
    doc.add_heading(t, level=1)


def h2(doc, t):
    doc.add_heading(t, level=2)


def para(doc, t, italic=False, muted=False, size=11):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.italic = italic
    r.font.size = Pt(size)
    if muted:
        r.font.color.rgb = MUTED
    return p


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        # **bold** qismlarini ajratib chiqarish
        parts = it.split("**")
        for i, seg in enumerate(parts):
            run = p.add_run(seg)
            run.bold = (i % 2 == 1)


def img(doc, name, width=6.2, caption=None):
    path = ASSETS / name
    if not path.exists():
        para(doc, f"[rasm topilmadi: {name}]", italic=True, muted=True)
        return
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(caption)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = MUTED


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, hh in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(hh)
        run.bold = True
        run.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(10)
    doc.add_paragraph()


def build():
    doc = Document()
    setup_styles(doc)

    # --- Sarlavha ---
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("MAMOGRAF — To'liq Foydalanuvchi Qo'llanmasi")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = ACCENT
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run("Mammografiya AI tahlil kompleksi (AI Scan)")
    rs.font.size = Pt(13)
    rs.font.color.rgb = MUTED
    d = doc.add_paragraph()
    d.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_run(date.today().strftime("%Y-%m-%d")).font.color.rgb = MUTED

    para(doc,
         "MAMOGRAF — mammografiya tasvirlarini AI yordamida tahlil qiluvchi, "
         "ko'krak saratonini erta aniqlashga mo'ljallangan dasturiy kompleks. "
         "DICOM ko'ruvchi, annotatsiya, AI inference, GPU'da model o'qitish, "
         "avtomatik hisobot generatsiyasi va PACS integratsiyasini birlashtiradi.")
    wp = doc.add_paragraph()
    wr = wp.add_run("⚠️ Klinik ogohlantirish: dastur tadqiqot va yordamchi maqsadida. "
                    "Barcha AI natijalari va hisobotlar radiolog tomonidan tasdiqlanishi shart.")
    wr.bold = True
    wr.font.color.rgb = RGBColor(0x9C, 0x27, 0x00)

    doc.add_page_break()

    # 1
    h1(doc, "1. Ishga tushirish")
    para(doc, "Dastur run.bat orqali ishga tushadi. U .venv muhitini yaratadi, "
              "GPU uchun CUDA torch (cu128) ni o'rnatadi va serverni "
              "http://127.0.0.1:8002 manzilida ishga tushiradi.")
    table(doc, ["Komponent", "Texnologiya"], [
        ["Backend", "FastAPI (Python)"],
        ["Frontend", "HTML / JavaScript / CSS"],
        ["DICOM", "pydicom, pylibjpeg"],
        ["AI", "Ultralytics YOLO, PyTorch (CUDA), radiomics"],
        ["Hisobot AI", "Lokal Ollama (qwen2.5) yoki shablon"],
        ["Ma'lumotlar bazasi", "SQLite"],
    ])

    # 2
    h1(doc, "2. Tizimga kirish va rollar")
    para(doc, "Birinchi ekran — tizimga kirish (foydalanuvchi nomi, parol, ixtiyoriy 2FA/TOTP).")
    img(doc, "ui_login.png", caption="Tizimga kirish ekrani")
    table(doc, ["Rol", "Imkoniyatlari"], [
        ["admin", "Hamma narsa: foydalanuvchilar, modellar, o'qitish, PACS, sozlamalar"],
        ["reviewer", "Annotatsiyalarni tekshirish/tasdiqlash, o'qitish, statistika, PACS"],
        ["annotator", "Tasvir ko'rish va annotatsiya (faqat o'ziniki)"],
    ])

    # 3
    h1(doc, "3. Asosiy ekran — DICOM ko'ruvchi")
    para(doc, "Chapda fayllar ro'yxati, markazda DICOM ko'ruvchi, o'ngda metadata, "
              "yuqorida asboblar paneli.")
    img(doc, "ui_viewer.png", caption="DICOM ko'ruvchi — belgilangan massa bilan")
    bullets(doc, [
        "**Yuklash** — DICOM yuklash, yuklashda avtomatik anonimlashtirish (PHI o'chiriladi).",
        "**Lokal disk** — server diskidagi DICOM'lar (LOCAL_DICOM_ROOT).",
        "**Window/Level** — yorqinlik/kontrast preset'lari.",
        "**Zoom / pan / 1:1 / Invert** — masshtab, surish, haqiqiy o'lcham, ranglarni teskari.",
        "**Multi-frame / 4-view** — ko'p kadr va 4 proeksiyani yonma-yon.",
        "**Heatmap** — AI ishonch issiqlik xaritasi.",
    ])

    # 4
    h1(doc, "4. Annotatsiya (lezyon belgilash)")
    bullets(doc, [
        "**Quti (bbox)** va **ko'pburchak (polygon)** — lezyon shaklini belgilash.",
        "**Ko'p yorliq (multi-label)** — bitta o'choqqa bir nechta sinf.",
        "**BI-RADS kategoriyasi** — har lezyonga (0–6, 4A/4B/4C).",
        "**Izoh (note)** — erkin matn.",
        "**Holat:** draft → submitted → approved / rejected (reviewer tasdig'i).",
    ])

    # 5
    h1(doc, "5. AI yordami (inference)")
    bullets(doc, [
        "**Avto-inference** — yuklashda avtomatik bbox'lar (ixtiyoriy).",
        "**Smart-click** — bir bosishda atrofdagi o'choqni topib quti chizadi.",
        "**Uncertainty heatmap** — model 'ishonchsiz' joylarini ko'rsatadi (TTA).",
        "**Batch inference** — bir nechta tasvirga bir vaqtda.",
        "**WBF ensemble + TTA** — bir nechta model/aylantirish natijasini birlashtirish.",
    ])
    para(doc, "Model arxitekturalari: BCA-YOLO, TILLNet-Det, XS-Classifier.", muted=True)

    # 6
    h1(doc, "6. Hisobot generatori")
    para(doc, "Belgilangan lezyonlardan mammografiya hisoboti qoralamasini avtomatik yaratadi "
              "(📝 Hisobot tugmasi).")
    img(doc, "ui_report.png", caption="Hisobot generatori — lokal AI (qwen2.5) bilan")
    table(doc, ["Usul", "Tavsif"], [
        ["Boyitilgan shablon", "Modelsiz, deterministik, bir zumda. Oflayn, eng xavfsiz."],
        ["Lokal AI (Ollama)", "Mashinangizdagi qwen2.5 tabiiy matn yozadi — kalit/internet shart emas."],
        ["Auto", "Lokal model bo'lsa undan, aks holda shablonga tushadi."],
    ])
    bullets(doc, [
        "Model **faqat berilgan topilmalardan** foydalanadi (to'qimaydi), temperature=0 (barqaror).",
        "Chiqish doim **QORALAMA** — radiolog tahrirlaydi va tasdiqlaydi.",
        "**💾 DICOM SR** — hisobotni standart DICOM Structured Report sifatida saqlash.",
        "**📋 Nusxa olish** — matnni clipboard'ga.",
    ])

    # 7
    h1(doc, "7. Model Studio — model o'qitish")
    para(doc, "O'z annotatsiyalaringizdan yangi YOLO modelini GPU'da o'qitish (🧪 Model Studio).")
    img(doc, "ui_studio.png", caption="Model Studio — GPU monitor va jonli grafiklar")
    bullets(doc, [
        "Dataset tanlash + data.yaml tekshiruvi (rasm soni, class taqsimoti, buzuq fayllar).",
        "Base model (YOLO11 n/s/m/l/x), pretrained, **Resume**.",
        "To'liq parametrlar: epochs, batch, imgsz, optimizer, LR, augmentation.",
        "**GPU monitor** — VRAM, utilization, harorat real vaqtda.",
        "**▶ Train / ⏹ Stop** — to'xtatilganda last.pt saqlanadi (Resume bilan davom).",
        "**Jonli grafiklar** — loss, mAP@50, mAP@50-95, precision/recall.",
        "**Avto-deploy** — o'qitilgan model app/models/ ga ko'chiriladi.",
    ])

    # 8
    h1(doc, "8. Dataset tayyorlash")
    para(doc, "Annotatsiyalardan Ultralytics YOLO datasetini avtomatik yaratadi (🎓 Dataset).")
    img(doc, "ui_datasetprep.png", caption="Dataset tayyorlash modali")
    bullets(doc, [
        "images/{train,val} + labels/{train,val} + data.yaml.",
        "**Bemor darajasida** train/val bo'linishi (data leakage'ning oldini oladi).",
        "Natija to'g'ridan-to'g'ri Model Studio'da ishlatishga tayyor.",
    ])

    # 9
    h1(doc, "9. Modellar dashboard")
    para(doc, "Mavjud AI modellarini boshqarish (🤖 Modellar).")
    img(doc, "ui_models.png", caption="Modellar dashboard")
    bullets(doc, [
        "Har model: hajm, qachon qo'shilgan, nechta annotation chiqargan.",
        "trained_* modellar — lokal o'qitilgan, o'chirish mumkin.",
        "**Active learning** — qayta o'qitish vaqtini tavsiya qiladi.",
    ])

    # 10
    h1(doc, "10. Eksport formatlari")
    table(doc, ["Format", "Tavsif"], [
        ["Annotated DICOM", "Annotatsiyalar bilan DICOM"],
        ["DICOM-SEG", "Segmentatsiya obyektlari (standart)"],
        ["DICOM-SR", "Strukturaviy hisobot (annotatsiya yoki to'liq hisobot)"],
        ["Mask (PNG)", "Bineriy/rangli maska"],
        ["COCO JSON", "Detection dataset formati"],
        ["CSV", "Annotatsiya tarixi / radiomics jadvallari"],
    ])

    # 11
    h1(doc, "11. Statistika")
    img(doc, "ui_stats.png", caption="Statistika paneli")
    bullets(doc, [
        "Jami DICOM, annotatsiyalar, bemorlar soni.",
        "Class taqsimoti, BI-RADS taqsimoti.",
        "Lokal DICOM ombori statistikasi.",
    ])

    # 12
    h1(doc, "12. PACS integratsiyasi")
    img(doc, "ui_pacs.png", caption="PACS serverlar")
    bullets(doc, [
        "PACS serverlarni qo'shish/boshqarish (AE Title, host, port).",
        "**C-ECHO** — ulanishni tekshirish.",
        "**C-FIND** — bemor/study qidirish.",
        "**C-STORE** — DICOM/SR ni PACS'ga jo'natish.",
        "**C-MOVE/GET** — PACS'dan tasvir olib kelish.",
        "**Worklist** — ish ro'yxatiga qo'shish.",
    ])

    # 13
    h1(doc, "13. Audit va tarix")
    img(doc, "ui_audit.png", caption="Audit timeline")
    bullets(doc, [
        "**Audit timeline** — kim, qachon, nimani o'zgartirdi.",
        "**Annotatsiya tarixi** — versiyalar (create/edit/approve/reject).",
        "**Bildirishnomalar** — tasdiqlash/rad etish xabarlari.",
    ])

    # 14
    h1(doc, "14. Foydalanuvchilar boshqaruvi")
    para(doc, "Faqat admin uchun (👥).")
    img(doc, "ui_admin.png", caption="Foydalanuvchilar boshqaruvi")
    bullets(doc, [
        "Foydalanuvchi qo'shish/tahrirlash/o'chirish, rol berish.",
        "Parolni qayta o'rnatish, faollashtirish/o'chirish.",
        "**2FA (TOTP)** — ikki bosqichli autentifikatsiya.",
    ])

    # 15
    h1(doc, "15. Xavfsizlik va maxfiylik")
    bullets(doc, [
        "**Avtomatik anonimlashtirish** — yuklashda PHI o'chiriladi; barcha qadamlar PHI-siz nusxada.",
        "**JWT autentifikatsiya** + rol asosidagi ruxsatlar.",
        "**TOTP 2FA** — ixtiyoriy ikki bosqichli kirish.",
        "**Lokal AI** — hisobotlar tashqi xizmatga jo'natilmaydi (Ollama mashinangizda).",
    ])

    # 16
    h1(doc, "16. Texnik arxitektura")
    img(doc, "architecture.png", caption="Tizim arxitekturasi")
    h2(doc, "Model arxitekturalari")
    img(doc, "arch_bca_yolo.png", width=5.6, caption="BCA-YOLO — ikki ko'krakni qiyoslab o'choq topish")
    img(doc, "arch_tillnet.png", width=5.6, caption="TILLNet-Det — rasm + klinik matn (multimodal)")
    img(doc, "cycle.png", width=5.0, caption="Active learning sikli (AI ↔ radiolog)")

    # 17
    h1(doc, "17. Tipik ish oqimi")
    for step in [
        "1. DICOM yuklash (avtomatik anonimlashtirish)",
        "2. AI yordami yoki qo'lda annotatsiya + BI-RADS",
        "3. Reviewer tasdiqlaydi (approved)",
        "4. 📝 Hisobot yaratish → tahrirlash → 💾 DICOM SR",
        "5. (Ixtiyoriy) 🎓 Dataset → 🧪 Model Studio → modelni o'qitish",
        "6. Yangi model bilan aniqlik oshadi → sikl takrorlanadi",
    ]:
        p = doc.add_paragraph(step)
        p.runs[0].font.size = Pt(11)

    doc.save(str(OUT))
    print("Saqlandi:", OUT)


if __name__ == "__main__":
    build()
